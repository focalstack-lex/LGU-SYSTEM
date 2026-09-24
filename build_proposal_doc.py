import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

def create_proposal_document():
    doc = docx.Document()

    # Define Color Palette (Executive Deep Navy & Maroon Accent)
    COLOR_PRIMARY = RGBColor(27, 54, 93)     # Deep Navy (#1B365D)
    HEX_PRIMARY = "1B365D"
    COLOR_SECONDARY = RGBColor(128, 0, 0)   # Deep Maroon (#800000)
    HEX_SECONDARY = "800000"
    COLOR_TEXT = RGBColor(34, 34, 34)        # Charcoal (#222222)
    COLOR_MUTED = RGBColor(90, 100, 110)     # Slate Grey (#5A646E)
    HEX_LIGHT_BG = "F4F6F9"
    HEX_BORDER = "CCCCCC"
    HEX_ROW_ALT = "F9FAFC"

    # Set Margins (0.55 in top/bottom, 0.75 in left/right)
    for section in doc.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        # Footer Setup
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("Financial Transparency & Student Academic Portal | 3-Page System Proposal")
        f_run.font.name = "Calibri"
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = COLOR_MUTED

    # Style Helper Functions
    def set_font(run, name="Calibri", size_pt=9.5, color=COLOR_TEXT, bold=False, italic=False):
        run.font.name = name
        run.font.size = Pt(size_pt)
        run.font.color.rgb = color
        run.bold = bold
        run.italic = italic

    def set_para_spacing(p, before=0, after=4, line_spacing=1.12):
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = line_spacing

    def set_cell_shading(cell, color_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_table_borders(table, color="CCCCCC"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_heading_1(text):
        p = doc.add_paragraph()
        set_para_spacing(p, before=9, after=4, line_spacing=1.12)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        set_font(run, name="Calibri", size_pt=12, color=COLOR_PRIMARY, bold=True)
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="2" w:color="{HEX_PRIMARY}"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)
        return p

    # --- PAGE 1: HEADER, METADATA, EXEC SUMMARY & CORE ARCHITECTURE ---
    
    # TOP HEADER BANNER IMAGE
    banner_path = os.path.join('client', 'assets', 'letterhead-banner.jpg')
    if os.path.exists(banner_path):
        p_banner = doc.add_paragraph()
        set_para_spacing(p_banner, before=0, after=4)
        p_banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_banner.add_run().add_picture(banner_path, width=Inches(7.0))

    # DOCUMENT TITLE & SUBTITLE
    p_title = doc.add_paragraph()
    set_para_spacing(p_title, before=1, after=1)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("SYSTEM PROPOSAL & INSTITUTIONAL BUDGET")
    set_font(r_title, name="Calibri", size_pt=16.5, color=COLOR_PRIMARY, bold=True)

    p_sub = doc.add_paragraph()
    set_para_spacing(p_sub, before=0, after=6)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Financial Transparency & Student Academic Portal")
    set_font(r_sub, name="Calibri", size_pt=11, color=COLOR_MUTED, bold=True)

    # METADATA SUMMARY BOX TABLE
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    set_table_borders(meta_table, HEX_PRIMARY)

    widths = [Inches(3.5), Inches(3.5)]
    meta_data = [
        [("PROJECT PROPONENT:", " College of Engineering LGU, Cor Jesu College"),
         ("SUBMITTED TO:", " Office of the President / Academic Administration")],
        [("DATE OF SUBMISSION:", " September 24, 2026 | Ref: SP-2026-COE-001"),
         ("PROPOSED BUDGET:", " ₱18,500.00 PHP (Turnkey Implementation)")]
    ]

    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = widths[col_idx]
            set_cell_shading(cell, HEX_LIGHT_BG)
            set_cell_margins(cell, top=65, bottom=65, left=100, right=100)
            lbl, val = meta_data[row_idx][col_idx]
            p = cell.paragraphs[0]
            set_para_spacing(p, before=0, after=0, line_spacing=1.05)
            r_l = p.add_run(lbl)
            set_font(r_l, size_pt=8.5, color=COLOR_PRIMARY, bold=True)
            r_v = p.add_run(val)
            set_font(r_v, size_pt=8.5, color=COLOR_TEXT, bold=False)

    # SECTION 1: EXECUTIVE SUMMARY
    add_heading_1("1. Executive Summary & Project Rationale")
    
    p_exec = doc.add_paragraph()
    set_para_spacing(p_exec, before=0, after=4, line_spacing=1.12)
    r = p_exec.add_run(
        "This proposal presents a turnkey deployment plan for the institutional adoption of the Financial Transparency and Student Academic Portal. Developed and validated within the College of Engineering Local Government Unit (COE-LGU) of Cor Jesu College, the platform resolves two critical administrative friction points: manual financial record-keeping and fragmented student academic progress verification."
    )
    set_font(r, size_pt=9.5)

    p_gaps = doc.add_paragraph()
    set_para_spacing(p_gaps, before=0, after=6, line_spacing=1.12)
    r = p_gaps.add_run(
        "Financial tracking relying on disconnected spreadsheets and paper receipts suffers from delayed liquidation reporting, lack of real-time balance visibility, and high audit overhead. Concurrently, academic evaluation (such as load approval, prerequisite validation, and standing verification) requires manual cross-referencing across separate rosters and forms. The proposed web-based system unifies both financial ledger management and academic prospectus verification into a single, secure, role-restricted platform."
    )
    set_font(r, size_pt=9.5)

    # SECTION 2: CORE SYSTEM CAPABILITIES
    add_heading_1("2. Core System Architecture & Scope")

    p_arch = doc.add_paragraph()
    set_para_spacing(p_arch, before=0, after=5, line_spacing=1.12)
    r = p_arch.add_run(
        "The application is built as a client-server web system leveraging Supabase (PostgreSQL with Row-Level Security) and standard web technologies. It requires zero client installation and functions seamlessly across desktop browsers, mobile devices (Progressive Web App), and optional native desktop clients."
    )
    set_font(r, size_pt=9.5)

    # Table 1: Core System Domains
    dom_table = doc.add_table(rows=5, cols=3)
    dom_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    dom_table.autofit = False
    set_table_borders(dom_table, HEX_BORDER)

    dom_widths = [Inches(1.5), Inches(4.3), Inches(1.2)]
    headers = ["Functional Domain", "Key Capabilities & Features", "Target Users"]
    
    hdr_row = dom_table.rows[0]
    for idx, cell in enumerate(hdr_row.cells):
        cell.width = dom_widths[idx]
        set_cell_shading(cell, HEX_PRIMARY)
        set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
        p = cell.paragraphs[0]
        set_para_spacing(p, before=0, after=0)
        r = p.add_run(headers[idx])
        set_font(r, size_pt=8.8, color=RGBColor(255, 255, 255), bold=True)

    domain_data = [
        ("1. Financial Ledger & Transparency",
         "Double-entry transaction ledger with compulsory receipt uploads, per-activity budget envelopes, real-time overdraft validation, automated PDF/Excel liquidation report generation, and public disclosure summary.",
         "Officers, Cashiers, Governor, Admin"),
        ("2. Academic Records & Load Verification",
         "Program prospectus structure, prerequisite and co-requisite validation, student load submission & Program Head approval workflow, credit-unit tracker, automated standing transcripts, and student roster integration.",
         "Students, Program Heads, Dean"),
        ("3. Student Services & Public Verification",
         "Targeted in-app & email notifications, structured student feedback facility, institutional announcements, and verifiable digital Curriculum Vitae (CV) builder with public token validation.",
         "Students, External Verifiers"),
        ("4. Executive Governance & Security",
         "8-tier Role-Based Access Control (RBAC), append-only immutable audit logging, institutional email domain lock (@cjc.edu.ph), rate limiting, row-level security policies, and weekly automated backups.",
         "System Administrator, Audit Officers")
    ]

    for row_idx, data in enumerate(domain_data, start=1):
        row = dom_table.rows[row_idx]
        bg_color = HEX_ROW_ALT if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = dom_widths[col_idx]
            set_cell_shading(cell, bg_color)
            set_cell_margins(cell, top=65, bottom=65, left=90, right=90)
            p = cell.paragraphs[0]
            set_para_spacing(p, before=0, after=0, line_spacing=1.08)
            r = p.add_run(text)
            is_bold = (col_idx == 0)
            set_font(r, size_pt=8.5, color=COLOR_PRIMARY if is_bold else COLOR_TEXT, bold=is_bold)

    # Page Break to start Page 2 crisply
    doc.add_page_break()

    # --- PAGE 2: SECURITY, ROADMAP & BUDGET PROPOSAL ---

    # SECTION 3: SECURITY & PRIVACY CONTROLS
    add_heading_1("3. Security, Privacy & Data Compliance")

    p_sec = doc.add_paragraph()
    set_para_spacing(p_sec, before=0, after=4, line_spacing=1.12)
    r = p_sec.add_run(
        "Security and privacy controls are embedded directly into the database architecture, complying strictly with the Data Privacy Act of 2012 (RA 10173):"
    )
    set_font(r, size_pt=9.5)

    sec_points = [
        ("Row-Level Security (RLS): ", "Database policies ensure students view only their own records, while officers access assigned organizational domains."),
        ("Append-Only Audit Logging: ", "Financial corrections or role alterations log actor ID, timestamp, and mandatory justification."),
        ("Domain-Locked Registration: ", "Registration is restricted to valid institutional email addresses (@cjc.edu.ph) via database triggers."),
        ("Encrypted Private Receipts: ", "Receipt images and document attachments are stored in private storage buckets served via authenticated tokens.")
    ]

    for title, desc in sec_points:
        p_pt = doc.add_paragraph()
        set_para_spacing(p_pt, before=0, after=3, line_spacing=1.08)
        p_pt.paragraph_format.left_indent = Inches(0.15)
        r1 = p_pt.add_run("• " + title)
        set_font(r1, size_pt=8.8, color=COLOR_PRIMARY, bold=True)
        r2 = p_pt.add_run(desc)
        set_font(r2, size_pt=8.8, color=COLOR_TEXT)

    # SECTION 4: IMPLEMENTATION ROADMAP
    add_heading_1("4. Phased Implementation Roadmap (4-Week Rollout)")

    p_road = doc.add_paragraph()
    set_para_spacing(p_road, before=0, after=4, line_spacing=1.12)
    r = p_road.add_run(
        "To guarantee an orderly rollout without disrupting routine academic operations, deployment is executed in four structured phases across four weeks:"
    )
    set_font(r, size_pt=9.5)

    pipe_table = doc.add_table(rows=5, cols=3)
    pipe_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    pipe_table.autofit = False
    set_table_borders(pipe_table, HEX_BORDER)

    pipe_widths = [Inches(1.3), Inches(4.2), Inches(1.5)]
    p_headers = ["Phase & Timeline", "Key Implementation Activities", "Primary Milestone"]

    hdr_row2 = pipe_table.rows[0]
    for idx, cell in enumerate(hdr_row2.cells):
        cell.width = pipe_widths[idx]
        set_cell_shading(cell, HEX_PRIMARY)
        set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
        p = cell.paragraphs[0]
        set_para_spacing(p, before=0, after=0)
        r = p.add_run(p_headers[idx])
        set_font(r, size_pt=8.8, color=RGBColor(255, 255, 255), bold=True)

    pipe_data = [
        ("Phase 1: Week 1\nDiscovery & Config",
         "Initial discovery, institution branding setup (logos, letterhead banners), database schema migration, email domain locking, and role permission mapping.",
         "Configured Environment"),
        ("Phase 2: Week 2\nData & Prospectus",
         "Encoding degree program prospectus structures, prerequisite rules, student roster CSV import, and initial general fund baseline setup.",
         "Populated Core Database"),
        ("Phase 3: Week 3\nPilot & Training",
         "Conducting pilot load evaluation with one department, user acceptance testing (UAT), training workshops for Program Heads, Cashiers, and Officers.",
         "Signed UAT & Readiness"),
        ("Phase 4: Week 4\nGo-Live & Support",
         "Full institutional production deployment, administrator console handover, SSL verification, automated backup activation, and 1-year support launch.",
         "Production Go-Live")
    ]

    for row_idx, data in enumerate(pipe_data, start=1):
        row = pipe_table.rows[row_idx]
        bg_color = HEX_ROW_ALT if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = pipe_widths[col_idx]
            set_cell_shading(cell, bg_color)
            set_cell_margins(cell, top=65, bottom=65, left=90, right=90)
            p = cell.paragraphs[0]
            set_para_spacing(p, before=0, after=0, line_spacing=1.08)
            r = p.add_run(text)
            is_bold = (col_idx == 0 or col_idx == 2)
            set_font(r, size_pt=8.5, color=COLOR_PRIMARY if is_bold else COLOR_TEXT, bold=is_bold)

    # SECTION 5: BUDGET PROPOSAL (15k - 20k)
    add_heading_1("5. Institutional Budget Proposal (₱18,500.00 Turnkey Package)")

    p_bud_intro = doc.add_paragraph()
    set_para_spacing(p_bud_intro, before=0, after=4, line_spacing=1.12)
    r = p_bud_intro.add_run(
        "Because the core system architecture is fully pre-built and operational, the adopting institution incurs ZERO software custom-development costs. The itemized budget below represents a turnkey deployment package spanning configuration, database migration, training, and a 1-year maintenance SLA."
    )
    set_font(r, size_pt=9.5)

    bud_table = doc.add_table(rows=6, cols=3)
    bud_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    bud_table.autofit = False
    set_table_borders(bud_table, HEX_PRIMARY)

    bud_widths = [Inches(1.5), Inches(4.2), Inches(1.3)]
    b_headers = ["Budget Component", "Scope of Deliverables & Services Included", "Cost (PHP)"]

    hdr_row3 = bud_table.rows[0]
    for idx, cell in enumerate(hdr_row3.cells):
        cell.width = bud_widths[idx]
        set_cell_shading(cell, HEX_PRIMARY)
        set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
        p = cell.paragraphs[0]
        if idx == 2:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_para_spacing(p, before=0, after=0)
        r = p.add_run(b_headers[idx])
        set_font(r, size_pt=8.8, color=RGBColor(255, 255, 255), bold=True)

    budget_items = [
        ("1. System Branding & Configuration",
         "Customization of institutional branding (school logo, seal, official letterhead banners), email domain lock setup (@cjc.edu.ph), and SMTP mail server integration.",
         "₱4,500.00"),
        ("2. Database Setup & Roster Migration",
         "Supabase PostgreSQL cloud instance provisioning, database schema migration execution (32 migration files), student roster CSV import, and curriculum prospectus encoding.",
         "₱5,000.00"),
        ("3. Institutional Training & Handover",
         "Role-based training workshops for Program Heads, Cashiers, Officers, and System Administrators. Provision of complete documentation manuals and administrator guides.",
         "₱4,000.00"),
        ("4. Security Hardening & 1-Yr Support",
         "SSL security configuration, automated weekly encrypted backup workflows, rate-limiting rules, priority bug fixes, and 12-month technical support agreement.",
         "₱5,000.00"),
        ("TOTAL INVESTMENT",
         "Complete, turnkey institutional deployment including configuration, training, documentation, and 1-year dedicated technical support. Zero recurring platform fees.",
         "₱18,500.00")
    ]

    for row_idx, data in enumerate(budget_items, start=1):
        row = bud_table.rows[row_idx]
        is_total = (row_idx == 5)
        bg_color = HEX_LIGHT_BG if is_total else (HEX_ROW_ALT if row_idx % 2 == 1 else "FFFFFF")
        
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = bud_widths[col_idx]
            set_cell_shading(cell, bg_color)
            set_cell_margins(cell, top=80 if is_total else 65, bottom=80 if is_total else 65, left=90, right=90)
            p = cell.paragraphs[0]
            if col_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_para_spacing(p, before=0, after=0, line_spacing=1.08)
            r = p.add_run(text)
            
            if is_total:
                set_font(r, size_pt=9.2, color=COLOR_PRIMARY if col_idx != 1 else COLOR_TEXT, bold=True)
            else:
                is_bold = (col_idx == 0)
                set_font(r, size_pt=8.5, color=COLOR_PRIMARY if is_bold else COLOR_TEXT, bold=is_bold)

    # Page Break to start Page 3 crisply
    doc.add_page_break()

    # --- PAGE 3: BENEFITS, GOVERNANCE & SIGN-OFF ---

    # SECTION 6: STRATEGIC BENEFITS & ROI MATRIX
    add_heading_1("6. Strategic Institutional Benefits & Expected Impact")

    p_ben = doc.add_paragraph()
    set_para_spacing(p_ben, before=0, after=6, line_spacing=1.12)
    r = p_ben.add_run(
        "The adoption of the system provides measurable, high-impact improvements across administrative governance, financial accountability, and academic operations:"
    )
    set_font(r, size_pt=9.5)

    benefits = [
        ("Immediate Liquidation Readiness: ", "Generates official PDF and Excel liquidation reports directly from recorded transaction ledgers, reducing report assembly time by over 90% and eliminating missing receipt bottlenecks."),
        ("Zero Manual Balance Discrepancies: ", "Real-time ledger balance validation prevents account overdrafts and maintains continuous transparency over organizational funds."),
        ("Automated Load & Prerequisite Evaluation: ", "Program Heads evaluate submitted loads against automated prerequisite rules, preventing invalid course sequence enrollments and manual transcript auditing."),
        ("Uncompromising Governance Audit Trail: ", "Append-only audit logs record every financial entry modification or administrative change, ensuring complete governance peace of mind and effortless internal audit preparation."),
        ("Omnichannel Student & Officer Access: ", "Students and faculty access real-time financial ledgers and academic standing from web browsers, installed mobile PWA apps, or native desktop applications.")
    ]

    for title, desc in benefits:
        p_b = doc.add_paragraph()
        set_para_spacing(p_b, before=0, after=4, line_spacing=1.1)
        p_b.paragraph_format.left_indent = Inches(0.15)
        r1 = p_b.add_run("• " + title)
        set_font(r1, size_pt=9, color=COLOR_SECONDARY, bold=True)
        r2 = p_b.add_run(desc)
        set_font(r2, size_pt=9, color=COLOR_TEXT)

    # SECTION 7: GOVERNANCE, DATA OWNERSHIP & SLA
    add_heading_1("7. Governance, Data Ownership & Technical Support SLA")

    p_gov = doc.add_paragraph()
    set_para_spacing(p_gov, before=0, after=5, line_spacing=1.12)
    r = p_gov.add_run(
        "Upon successful deployment, the adopting institution assumes full, unencumbered ownership of all database records, user credentials, and financial documentation. System database backups are encrypted weekly and stored within institutional storage repositories. A formal Data Processing Agreement (DPA) will be executed alongside system handover to solidify data compliance."
    )
    set_font(r, size_pt=9.5)

    p_sla = doc.add_paragraph()
    set_para_spacing(p_sla, before=0, after=8, line_spacing=1.12)
    r = p_sla.add_run(
        "The turnkey package includes a 12-Month Maintenance & Support SLA covering system updates, SSL maintenance, database optimizations, and priority technical assistance for System Administrators."
    )
    set_font(r, size_pt=9.5)

    # SECTION 8: RECOMMENDATION & SIGN-OFF
    add_heading_1("8. Recommendation & Authorization Sign-Off")

    p_rec = doc.add_paragraph()
    set_para_spacing(p_rec, before=0, after=18, line_spacing=1.12)
    r = p_rec.add_run(
        "We respectfully request formal authorization to proceed with Phase 1 (Discovery & Institutional Configuration) under the proposed turnkey budget of ₱18,500.00 PHP."
    )
    set_font(r, size_pt=9.5)

    # Signature Block Table
    sig_table = doc.add_table(rows=2, cols=3)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    set_table_borders(sig_table, "FFFFFF")

    sig_widths = [Inches(2.3), Inches(2.3), Inches(2.4)]
    sig_data = [
        [("PREPARED BY:", "\n\n_______________________\nLead Systems Proponent\nCollege of Engineering LGU"),
         ("NOTED BY:", "\n\n_______________________\nFaculty Adviser / Dept. Head\nCollege of Engineering"),
         ("APPROVED BY:", "\n\n_______________________\nSchool Head / President\nAdopting Institution")],
        [("Date: _______________", ""), ("Date: _______________", ""), ("Date: _______________", "")]
    ]

    for row_idx, row in enumerate(sig_table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = sig_widths[col_idx]
            set_cell_margins(cell, top=40, bottom=40, left=40, right=40)
            p = cell.paragraphs[0]
            set_para_spacing(p, before=0, after=0, line_spacing=1.1)
            hdr_txt, body_txt = sig_data[row_idx][col_idx]
            
            r_h = p.add_run(hdr_txt)
            set_font(r_h, size_pt=8.5, color=COLOR_PRIMARY, bold=True)
            if body_txt:
                r_b = p.add_run(body_txt)
                set_font(r_b, size_pt=8.5, color=COLOR_TEXT, bold=False)

    doc_path = "System Proposal - Financial Transparency and Student Academic Portal.docx"
    doc.save(doc_path)
    print(f"Updated document saved to: {doc_path}")

if __name__ == "__main__":
    create_proposal_document()
